"""v0.3 RAG 引擎单元测试（§18.1）。

与 test_identity.py 同一约定：CI 环境无 Postgres/Redis/网络/torch，
全部为内存即时构造的纯逻辑测试（解析器 fixture 由 python-docx/python-pptx/PyMuPDF 内存生成）。
"""

import io
from types import SimpleNamespace

import fitz  # PyMuPDF
import pytest
from docx import Document as DocxDocument
from pptx import Presentation

from app.core.chunking import semantic_splitter
from app.core.chunking.semantic_splitter import (
    TARGET_MAX,
    make_chunk_id,
    split_blocks,
)
from app.core.parser import parse_by_file_type
from app.core.parser.base import ParsedBlock, ParseError
from app.core.parser.md_parser import MarkdownParser
from app.core.parser.pdf_parser import PdfParser
from app.core.parser.ppt_parser import PptParser
from app.core.parser.word_parser import WordParser
from app.core.pipeline import extract_citations
from app.core.retrieval.fine import bm25_search, tokenize
from app.core.retrieval.fusion import rrf_fuse
from app.workers.parse_worker import build_chunk_meta

# ---------------------------------------------------------------------------
# Markdown 解析器
# ---------------------------------------------------------------------------


def test_md_parser_heading_hierarchy():
    md = (
        "# 第一章 绪论\n\n第一节正文。\n\n## 1.1 背景\n\n小节正文。\n\n"
        "# 第二章\n\n第二章正文。\n"
    ).encode()
    blocks = MarkdownParser().parse(md)
    titles = [(b.title, b.level) for b in blocks if b.level >= 1]
    assert titles == [("第一章 绪论", 1), ("1.1 背景", 2), ("第二章", 1)]
    contents = [b.content for b in blocks if b.level == 0]
    assert contents == ["第一节正文。", "小节正文。", "第二章正文。"]


def test_md_parser_no_headings():
    blocks = MarkdownParser().parse("纯正文\n没有标题".encode())
    assert len(blocks) == 1
    assert blocks[0].level == 0 and blocks[0].title is None


# ---------------------------------------------------------------------------
# Word 解析器
# ---------------------------------------------------------------------------


def _make_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("第一章 操作系统概述", level=1)
    doc.add_paragraph("操作系统是管理计算机硬件与软件资源的程序。")
    doc.add_heading("1.1 操作系统的目标", level=2)
    doc.add_paragraph("方便性、有效性、可扩充性和开放性。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_word_parser_styles():
    blocks = WordParser().parse(_make_docx())
    assert [(b.title, b.level) for b in blocks if b.level >= 1] == [
        ("第一章 操作系统概述", 1),
        ("1.1 操作系统的目标", 2),
    ]
    contents = [b.content for b in blocks if b.level == 0]
    assert contents[0].startswith("操作系统是管理")
    assert contents[1].startswith("方便性")


# ---------------------------------------------------------------------------
# PPT 解析器
# ---------------------------------------------------------------------------


def _make_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = "进程与线程"
    slide.placeholders[1].text_frame.text = "进程是资源分配的最小单位。"
    slide.notes_slide.notes_text_frame.text = "强调与线程的区别"
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "调度算法"
    slide2.placeholders[1].text_frame.text = "先来先服务、短作业优先。"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_ppt_parser_pages():
    blocks = PptParser().parse(_make_pptx())
    assert len(blocks) == 2
    first = blocks[0]
    assert first.title == "进程与线程"
    assert first.level == 1 and first.page_num == 1
    assert "进程是资源分配" in first.content
    assert "备注：强调与线程的区别" in first.content
    assert blocks[1].page_num == 2


# ---------------------------------------------------------------------------
# PDF 解析器
# ---------------------------------------------------------------------------


def _make_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "第一章 内存管理", fontsize=20, fontname="china-s")
    page.insert_text(
        (72, 110),
        "分页存储管理的基本原理是把逻辑地址空间划分为若干固定大小的页。",
        fontsize=11,
        fontname="china-s",
    )
    page2 = doc.new_page()
    page2.insert_text((72, 72), "3.1 分页存储管理", fontsize=14, fontname="china-s")
    page2.insert_text(
        (72, 110),
        "页表的作用是实现从页号到物理块号的地址映射，这是考试常考的知识点。",
        fontsize=11,
        fontname="china-s",
    )
    data = doc.tobytes()
    doc.close()
    return data


def test_pdf_parser_titles_by_font_size():
    blocks = PdfParser().parse(_make_pdf())
    titles = [(b.title, b.level, b.page_num) for b in blocks if b.level >= 1]
    assert ("第一章 内存管理", 1, 1) in titles
    assert ("3.1 分页存储管理", 2, 2) in titles
    body = [b for b in blocks if b.level == 0]
    assert all(b.page_num in (1, 2) for b in body)
    assert any("逻辑地址空间" in b.content for b in body)


def test_pdf_parser_scanned_raises():
    doc = fitz.open()
    doc.new_page()  # 空白页，模拟扫描件无文字层
    data = doc.tobytes()
    doc.close()
    with pytest.raises(ParseError, match="扫描件"):
        PdfParser().parse(data)


def test_parse_by_file_type_dispatch():
    blocks = parse_by_file_type("markdown", "# 标题\n\n正文".encode())
    assert blocks[0].title == "标题"
    with pytest.raises(ParseError, match="不支持的文件类型"):
        parse_by_file_type("epub", b"x")


# ---------------------------------------------------------------------------
# 语义切块
# ---------------------------------------------------------------------------


def test_split_blocks_respects_heading_boundaries():
    blocks = [
        ParsedBlock(title="第一章", content="", level=1),
        ParsedBlock(title=None, content="第一章的第一段正文内容，足够长以不被合并。" * 3, level=0),
        ParsedBlock(title="第二章", content="", level=1),
        ParsedBlock(title=None, content="第二章的正文内容，同样足够长以不被合并掉。" * 3, level=0),
    ]
    chapter_ids = [101, 101, 102, 102]
    chunks = split_blocks(blocks, document_id=7, scope="personal", chapter_ids=chapter_ids)
    assert len(chunks) == 2
    assert chunks[0]["chapter_id"] == 101 and chunks[1]["chapter_id"] == 102
    assert chunks[0]["chunk_id"] == "per_d7_00001"
    assert chunks[1]["chunk_id"] == "per_d7_00002"


def test_split_blocks_long_text_with_overlap():
    sentence = "这是一个关于操作系统分页机制的详细描述句子，用来凑够切块长度。"
    blocks = [ParsedBlock(title=None, content=sentence * 40, level=0, page_num=3)]
    chunks = split_blocks(blocks, document_id=1, scope="public")
    assert len(chunks) >= 2
    assert all(len(c["content"]) <= TARGET_MAX + 60 for c in chunks)
    # 相邻块存在 overlap：后块开头内容来自前块结尾
    assert chunks[1]["content"][:20] in chunks[0]["content"]
    assert all(c["page_num"] == 3 for c in chunks)
    # public scope 的 chunk_id 前缀
    assert chunks[0]["chunk_id"].startswith("pub_d1_")


def test_split_blocks_chunk_id_unique_and_idempotent():
    blocks = [
        ParsedBlock(title="第一章", content="", level=1),
        ParsedBlock(title=None, content="短段落一。" * 30, level=0),
        ParsedBlock(title=None, content="短段落二。" * 30, level=0),
    ]
    first = split_blocks(blocks, document_id=9, scope="personal")
    second = split_blocks(blocks, document_id=9, scope="personal")
    assert [c["chunk_id"] for c in first] == [c["chunk_id"] for c in second]
    assert [c["content"] for c in first] == [c["content"] for c in second]
    assert len({c["chunk_id"] for c in first}) == len(first)


def test_split_blocks_merges_short_adjacent():
    """同属一章的相邻过短块合并为一个（跨章不合并）。"""
    blocks = [
        ParsedBlock(title="第一章", content="", level=1),
        ParsedBlock(title=None, content="很短。", level=0),
        ParsedBlock(title="第二章", content="", level=1),
        ParsedBlock(title=None, content="也很短。", level=0),
    ]
    # 不传 chapter_ids：两章均映射为 None（同章节键），触发合并
    chunks = split_blocks(blocks, document_id=2, scope="personal")
    assert len(chunks) == 1
    assert "很短" in chunks[0]["content"] and "也很短" in chunks[0]["content"]
    # 章节 ID 不同时不合并
    separate = split_blocks(
        blocks, document_id=2, scope="personal", chapter_ids=[1, 1, 2, 2]
    )
    assert len(separate) == 2


def test_make_chunk_id_format():
    assert make_chunk_id("personal", 12, 3) == "per_d12_00003"
    assert make_chunk_id("public", 12, 3) == "pub_d12_00003"


def test_split_blocks_section_from_level2():
    blocks = [
        ParsedBlock(title="第一章", content="", level=1),
        ParsedBlock(title="1.1 背景", content="", level=2),
        ParsedBlock(title=None, content="小节正文，足够长以独立成块不被合并。" * 3, level=0),
    ]
    chunks = split_blocks(blocks, document_id=3, scope="personal")
    assert chunks[0]["section"] == "1.1 背景"


# ---------------------------------------------------------------------------
# RRF 融合与 BM25 分词
# ---------------------------------------------------------------------------


def _hit(chunk_id: str):
    return SimpleNamespace(chunk_id=chunk_id)


def test_rrf_fuse_orders_and_weights():
    semantic = [_hit("a"), _hit("b"), _hit("c")]
    bm25 = [_hit("b"), _hit("d")]
    fused = rrf_fuse(semantic, bm25, top_k=3)
    ids = [chunk_id for chunk_id, _ in fused]
    assert ids[0] == "b"  # 两路都命中，分数最高
    assert set(ids) == {"a", "b", "c"}
    scores = dict(fused)
    assert scores["b"] > scores["a"] > scores["c"]


def test_rrf_fuse_weighted_personal_boost():
    """mixed 模式个人库加权 1.2：同 rank 的个人块应压过公共块。"""
    personal = [_hit("per_d1_00001")]
    public = [_hit("pub_d1_00001")]
    fused = rrf_fuse(personal, public, weight_semantic=1.2, weight_bm25=1.0, top_k=2)
    assert fused[0][0] == "per_d1_00001"


def test_tokenize_bigram_and_words():
    tokens = tokenize("分页storage管理")
    assert "分页" in tokens and "管理" in tokens
    assert "storage" in tokens
    assert tokenize("") == []


async def test_bm25_search_ranks_relevant_first():
    chunks = [
        SimpleNamespace(chunk_id="c1", content="分页存储管理把逻辑地址空间划分为固定大小的页"),
        SimpleNamespace(chunk_id="c2", content="编译原理中的词法分析器负责识别单词符号"),
        SimpleNamespace(chunk_id="c3", content="分页与分段的区别在于页是定长的物理划分"),
    ]
    hits = await bm25_search("分页存储", chunks, top_k=2)
    assert hits and hits[0].chunk_id in ("c1", "c3")
    assert all(h.chunk_id != "c2" for h in hits)


async def test_bm25_search_empty_corpus():
    assert await bm25_search("任意", [], top_k=5) == []


# ---------------------------------------------------------------------------
# 引用提取与 chunk 元数据组装
# ---------------------------------------------------------------------------


def test_extract_citations_filters_and_dedupes():
    answer = "分页是定长划分 [per_d1_00001]，分段按逻辑单位 [pub_d2_00003]。" \
        "再次出现 [per_d1_00001]，以及幻觉引用 [per_d9_99999]。"
    valid = {"per_d1_00001", "pub_d2_00003"}
    assert extract_citations(answer, valid) == ["per_d1_00001", "pub_d2_00003"]
    assert extract_citations("没有引用的回答", valid) == []


def test_build_chunk_meta_nullable_sentinels():
    doc = SimpleNamespace(id=5, course_id=12, owner_id=34, file_type="pdf_textbook")
    meta = build_chunk_meta(
        doc, {"chunk_id": "per_d5_00001", "chapter_id": None, "page_num": None}
    )
    assert meta == {
        "chunk_id": "per_d5_00001",
        "document_id": 5,
        "course_id": 12,
        "owner_id": 34,
        "chapter_id": -1,
        "page_num": -1,
        "file_type": "pdf_textbook",
    }
    assert all(isinstance(v, str | int) for v in meta.values())  # Chroma 仅接受标量
    meta2 = build_chunk_meta(
        doc, {"chunk_id": "per_d5_00002", "chapter_id": 7, "page_num": 78}
    )
    assert meta2["chapter_id"] == 7 and meta2["page_num"] == 78


def test_semantic_splitter_constants_sane():
    assert semantic_splitter.TARGET_MIN < semantic_splitter.TARGET_MAX
    assert semantic_splitter.OVERLAP < semantic_splitter.TARGET_MIN
